"""Stage 10, Excel export: the sheet's JSON (fpx.export.to_json) as a workbook for area management.

Sheets: Räume (one row per room with the three areas and the review state), Geschoss (floor figures), Öffnungen,
QA, Meta. Numbers are numbers, the header row is frozen and filterable.
"""
from .config import DEFAULT

ROOMS = [("ID", "id"), ("AOID", "aoid"), ("Raumnummer", "number"), ("Name", "name"), ("Nutzung", "usage"),
         ("Netto m²", "area_net"), ("Brutto m² (Wandanteil)", "area_gross"), ("Polygon m²", "area_polygon"), ("Stempel m²", "area_stamp"),
         ("Abweichung %", "area_deviation_pct"), ("Vergleichsbasis", "area_basis"), ("Konfidenz", "confidence"), ("Hinweise", "reasons"),
         ("Treppenläufe", "stair_flights"), ("Nachbarn", "neighbours"), ("AOIDs gelesen", "aoids_read"), ("Name roh", "name_raw"),
         ("Kleinfläche", "small_region")]
OPENINGS = [("ID", "id"), ("Art", "kind"), ("Wand", "host"), ("Aussen", "exterior"), ("Breite m", "width"), ("Verbindet", "connects"), ("Quelle", "source")]
QA = [("Prüfung", "check"), ("Schwere", "severity"), ("Element", "element"), ("Meldung", "message")]


def _cell(v):
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v if x is not None) if v else None
    if isinstance(v, dict):
        return str(v)
    return v


def to_xlsx(data, path, cfg=DEFAULT):
    """Write the sheet's JSON to path as .xlsx. Returns the row counts per sheet."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    counts = {}

    def table(title, columns, rows, widths=None):
        ws = wb.create_sheet(title) if wb.active.title != "Sheet" else wb.active
        ws.title = title
        ws.append([c for c, _ in columns])
        for r in rows:
            ws.append([_cell(r.get(k)) for _, k in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 2)}"
        for i, (c, _) in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(10, min(40, len(c) + 4)))
        counts[title] = len(rows)
        return ws

    table("Räume", ROOMS, data["rooms"], {"Name": 24, "Hinweise": 60, "Nachbarn": 20})
    fl = data["floor"]
    floor_rows = [{"k": "GF m² (Geschossfläche)", "v": fl.get("gf_area")}, {"k": "AGF m² (Aussengeschossfläche, Balkone)", "v": fl.get("agf_area")},
                  {"k": "Summe Netto m²", "v": fl.get("sum_net")}, {"k": "Summe Brutto m² (Wandanteile)", "v": fl.get("sum_gross")},
                  {"k": "EBF-Vorschlag m²", "v": (fl.get("ebf_proposal") or {}).get("area")},
                  {"k": "EBF ausgeschlossen", "v": (fl.get("ebf_proposal") or {}).get("excluded")},
                  {"k": "Räume", "v": len(data["rooms"])}, {"k": "Öffnungen", "v": len(data["openings"])},
                  {"k": "Lufträume / Treppenaugen abgezogen", "v": [v.get("label") for v in data.get("voids", []) if v.get("gf_deducted")]},
                  {"k": "Wandstücke", "v": len(data.get("walls", []))}, {"k": "Geschlossene Lücken (Stufe 3b)", "v": len(data.get("wall_bridges", []))}]
    table("Geschoss", [("Kennzahl", "k"), ("Wert", "v")], floor_rows, {"Kennzahl": 40, "Wert": 30})
    table("Öffnungen", OPENINGS, data["openings"])
    table("QA", QA, data.get("qa", []), {"Meldung": 90, "Prüfung": 26})
    sh, sc = data["sheet"], data["sheet"].get("scale") or {}
    meta = [{"k": "Schema", "v": data.get("schema_version")}, {"k": "Generator", "v": (data.get("generator") or {}).get("name")},
            {"k": "Blatt", "v": sh.get("id")}, {"k": "Titel", "v": sh.get("title")}, {"k": "Quelle", "v": sh.get("source")},
            {"k": "Eingabeklasse", "v": sh.get("input_class")}, {"k": "Massstab", "v": sc.get("value")}, {"k": "Massstab Methode", "v": sc.get("method")},
            {"k": "Hinweis", "v": (data.get("generator") or {}).get("note")}]
    table("Meta", [("Feld", "k"), ("Wert", "v")], meta, {"Feld": 20, "Wert": 80})
    wb.save(str(path))
    return counts
